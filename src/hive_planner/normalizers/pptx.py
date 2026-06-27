from __future__ import annotations

from pathlib import Path

from hive_planner.models import NormalizeResult
from hive_planner.normalizers import EmptySourceError, MissingDependencyError


def _slide_text(slide) -> str:
    parts: list[str] = []
    if slide.shapes.title and slide.shapes.title.text:
        parts.append(slide.shapes.title.text.strip())
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        text = shape.text.strip()
        if not text:
            continue
        if slide.shapes.title and text == slide.shapes.title.text.strip():
            continue
        parts.append(text)
    return "\n".join(parts).strip()


def normalize_pptx(file_path: Path) -> NormalizeResult:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise MissingDependencyError(
            "PPTX support requires python-pptx: pip install -e '.[pptx]'"
        ) from exc

    title = file_path.stem.replace("-", " ").title()
    try:
        presentation = Presentation(file_path)
    except Exception as exc:
        raise EmptySourceError(f"Could not open PPTX: {exc}") from exc

    lines: list[str] = [f"# {title}", ""]
    raw_parts: list[str] = []
    slide_num = 0

    for slide in presentation.slides:
        slide_num += 1
        text = _slide_text(slide)
        if not text:
            continue
        lines.extend([f"## Slide {slide_num}", "", text, ""])
        raw_parts.append(text)

    if not raw_parts:
        raise EmptySourceError("PPTX has no extractable text")

    normalized_body = "\n".join(lines).strip() + "\n"
    return NormalizeResult(
        normalized_body=normalized_body,
        raw_text="\n\n".join(raw_parts),
        title=title,
        extra_metadata={"format": "pptx", "slide_count": slide_num},
    )
