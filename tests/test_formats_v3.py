from __future__ import annotations

from pathlib import Path

import pytest

from hive_planner.chunker import chunk_markdown
from hive_planner.config import RepoPaths
from hive_planner.indexer import load_chunk_index, load_document_index
from hive_planner.normalizer import normalize_document
from hive_planner.normalizers import MissingDependencyError
from hive_planner.normalizers.html import normalize_html
from hive_planner.pipeline import ingest


def _make_html(path: Path, body: str, title: str = "Sample Page") -> None:
    path.write_text(
        f"<!DOCTYPE html><html><head><title>{title}</title></head><body>{body}</body></html>",
        encoding="utf-8",
    )


def _make_pptx(path: Path, slides: list[str]) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for text in slides:
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = text
        slide.placeholders[1].text = f"Body for {text}"
    presentation.save(path)


@pytest.fixture
def html_dep():
    pytest.importorskip("bs4")
    return True


@pytest.fixture
def pptx_dep():
    pytest.importorskip("pptx")
    return True


def test_empty_html_skipped(html_dep, tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.ensure_dirs()
    empty = paths.sources / "empty.html"
    empty.write_text("<html><head><title>Empty</title></head><body></body></html>", encoding="utf-8")
    report = ingest(paths)
    assert report.skipped == 1
    assert any("extractable" in w.lower() or "empty" in w.lower() for w in report.warnings)


def test_corrupt_html_skipped(html_dep, tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    bad = paths.sources / "bad.html"
    bad.write_bytes(b"\xff\xfe not html")
    (paths.sources / "ok.md").write_text("# OK\n\nstill works\n", encoding="utf-8")
    report = ingest(paths)
    assert report.processed == 1
    assert report.skipped == 1


def test_missing_html_dep_warning(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    html = paths.sources / "need.html"
    html.write_text("<html><body>test</body></html>", encoding="utf-8")

    def fail_html(_path: Path):
        raise MissingDependencyError("HTML support requires beautifulsoup4")

    monkeypatch.setattr("hive_planner.normalizer.normalize_html", fail_html)
    (paths.sources / "ok.md").write_text("# OK\n\nmd still works\n", encoding="utf-8")
    report = ingest(paths)
    assert report.processed == 1
    assert report.skipped == 1
    assert any("beautifulsoup4" in w for w in report.warnings)


def test_csv_without_headers(tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    csv_path = paths.sources / "notes.csv"
    csv_path.write_text("line one\nline two\nline three\n", encoding="utf-8")
    report = ingest(paths)
    assert report.processed == 1
    chunks = load_chunk_index(paths.chunk_index).chunks
    assert chunks
    assert any("line one" in c.text for c in chunks)


def test_html_reingest_same_doc_id(html_dep, tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    html = paths.sources / "ref" / "doc.html"
    html.parent.mkdir(parents=True)
    _make_html(html, "<h2>Version One</h2><p>first content</p>")
    ingest(paths)
    doc_id_1 = load_document_index(paths.document_index).documents[0].doc_id
    hash_1 = load_document_index(paths.document_index).documents[0].content_hash

    _make_html(html, "<h2>Version Two</h2><p>changed content</p>")
    ingest(paths)
    doc = load_document_index(paths.document_index).documents[0]
    assert doc.doc_id == doc_id_1
    assert doc.content_hash != hash_1


def test_html_skips_duplicate_title_h1(html_dep, tmp_path: Path) -> None:
    html_path = tmp_path / "page.html"
    html_path.write_text(
        "<html><head><title>My Page</title></head>"
        "<body><h1>My Page</h1><p>Body text here.</p></body></html>",
        encoding="utf-8",
    )
    result = normalize_html(html_path)
    assert result.normalized_body.count("My Page") == 1
    assert "Body text here" in result.normalized_body


def test_html_normalization(html_dep, tmp_path: Path) -> None:
    html_path = tmp_path / "page.html"
    _make_html(html_path, "<h2>Section</h2><p>Hello HTML world.</p>", title="My Page")
    result = normalize_html(html_path)
    assert "My Page" in result.normalized_body
    assert "Section" in result.normalized_body
    chunks = chunk_markdown("doc1", result.normalized_body)
    assert any("Hello" in c.text for c in chunks)


def test_pptx_slide_metadata(pptx_dep, tmp_path: Path) -> None:
    pptx_path = tmp_path / "deck.pptx"
    _make_pptx(pptx_path, ["Intro", "Details"])
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    pptx_path = paths.sources / "deck.pptx"
    _make_pptx(pptx_path, ["Intro", "Details"])
    report = ingest(paths)
    assert report.processed == 1
    chunks = load_chunk_index(paths.chunk_index).chunks
    pages = [c.page for c in chunks if c.page is not None]
    assert 1 in pages or 2 in pages


def test_csv_table_format(tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    csv_path = paths.sources / "data.csv"
    csv_path.write_text("name,value\nalpha,1\nbeta,2\n", encoding="utf-8")
    report = ingest(paths)
    assert report.processed == 1
    doc, body, _ = normalize_document(paths, csv_path)
    assert "| name | value |" in body
    assert doc.format == "csv"


def test_csv_escapes_pipe_in_cells(tmp_path: Path) -> None:
    paths = RepoPaths(root=tmp_path)
    paths.sources.mkdir(parents=True)
    csv_path = paths.sources / "pipes.csv"
    csv_path.write_text("name,note\nalpha,a|b\n", encoding="utf-8")
    _, body, _ = normalize_document(paths, csv_path)
    assert "a\\|b" in body
